from pptx import Presentation
from pptx.util import Pt

import config


class PPTBuilder:

    def __init__(self):

        self.prs = Presentation()

        self.prs.slide_width = (
            config.SLIDE_WIDTH
        )

        self.prs.slide_height = (
            config.SLIDE_HEIGHT
        )

        self.layout = (
            self.prs.slide_layouts[6]
        )

    def add_slide(
        self,
        result
    ):

        slide = (
            self.prs.slides.add_slide(
                self.layout
            )
        )

        title_box = (
            slide.shapes.add_textbox(
                config.TITLE_X,
                config.TITLE_Y,
                config.TITLE_W,
                config.TITLE_H
            )
        )

        title_box.text_frame.text = (
            result["title"]
        )

        title_box.text_frame.paragraphs[
            0
        ].font.size = Pt(
            config.TITLE_FONT_SIZE
        )

        slide.shapes.add_picture(
            result["model_file"],
            config.MODEL_X,
            config.MODEL_Y,
            width=config.MODEL_WIDTH
        )

        if result["legend_file"]:

            slide.shapes.add_picture(
                result["legend_file"],
                config.LEGEND_X,
                config.LEGEND_Y,
                width=config.LEGEND_WIDTH
            )

    def save(self):

        self.prs.save(
            config.OUTPUT_PPT
        )