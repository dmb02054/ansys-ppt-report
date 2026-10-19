"""편집 가능한 PPT 템플릿에 처리된 이미지와 Case 데이터를 채워 넣는 모듈.

템플릿의 결과 슬라이드·정리 슬라이드를 복제해 내용을 채운 뒤,
저장 시 편집용 템플릿 원본 슬라이드는 제거한다.
"""

import re
from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from .. import config


class PPTBuilder:

    def __init__(self):

        self.prs = Presentation(
            str(
                config.TEMPLATE_FILE
            )
        )

        self.result_template_index = (
            self.find_result_template()
        )
        self.summary_template_indices = (
            self.find_summary_templates()
        )
        template_indices = {
            self.result_template_index,
            *self.summary_template_indices
        }
        self.base_slide_count = (
            len(self.prs.slides) - len(template_indices)
        )
        self.generated_case_count = 0

    # =====================================
    # 결과 템플릿 찾기
    # =====================================

    def find_result_template(self):

        """템플릿에서 결과 슬라이드를 찾는다.

        [TEMPLATE_RESULT] 태그를 먼저 찾고, 사용자가 태그를 지웠다면
        이미지 자리표시자 태그가 모두 있는 슬라이드를 대신 사용한다.
        """

        fallback_index = None
        required_tags = {
            config.TITLE_TAG,
            "[RACK_MODEL]",
            "[RACK_LEGEND]",
            "[GLASS_MODEL]",
            "[GLASS_LEGEND]"
        }

        for idx, slide in enumerate(
            self.prs.slides
        ):

            slide_text = []

            for shape in slide.shapes:

                if getattr(shape, "has_table", False):
                    slide_text.extend(
                        cell.text
                        for row in shape.table.rows
                        for cell in row.cells
                    )

                if not hasattr(
                    shape,
                    "text"
                ):
                    continue

                slide_text.append(shape.text)

                if (
                    config.RESULT_TEMPLATE_TAG
                    in shape.text
                ):

                    return idx

            combined_text = "\n".join(slide_text)
            if all(tag in combined_text for tag in required_tags):
                fallback_index = idx

        if fallback_index is not None:
            return fallback_index

        raise ValueError(
            f"{config.RESULT_TEMPLATE_TAG}"
            f" 슬라이드를 찾을 수 없습니다."
        )

    def find_summary_templates(self):

        """'구분/변위' 헤더와 Case 행을 가진 정리 예시 슬라이드 목록을 찾는다."""

        summary_indices = []
        for index, slide in enumerate(self.prs.slides):
            if index == self.result_template_index:
                continue

            for shape in slide.shapes:
                if not getattr(shape, "has_table", False):
                    continue

                table = shape.table
                table_text = {
                    cell.text.strip()
                    for row in table.rows
                    for cell in row.cells
                    if cell.text.strip()
                }
                has_case_rows = any(
                    re.match(r"Case\s*\d+", text, re.IGNORECASE)
                    for text in table_text
                )
                if {"구분", "변위"}.issubset(table_text) and has_case_rows:
                    summary_indices.append(index)
                    break

        return summary_indices

    # =====================================
    # 슬라이드 복제
    # =====================================

    def duplicate_slide(
        self,
        slide_index
    ):

        """템플릿 슬라이드를 도형·서식째 깊은 복사로 새 슬라이드에 복제한다."""

        source = (
            self.prs.slides[
                slide_index
            ]
        )

        new_slide = (
            self.prs.slides.add_slide(
                source.slide_layout
            )
        )

        for shape in source.shapes:

            new_slide.shapes._spTree.insert_element_before(
                deepcopy(
                    shape.element
                ),
                "p:extLst"
            )

        return new_slide

    # =====================================
    # 이미지 자리표시자 교체
    # =====================================

    def replace_placeholder_image(
        self,
        slide,
        placeholder,
        image_file
    ):

        """자리표시자 도형을 지우고 그 자리에 이미지를 비율 유지로 가운데 배치한다."""

        if not image_file:
            return

        for shape in slide.shapes:

            if not hasattr(
                shape,
                "text"
            ):
                continue

            if (
                shape.text.strip()
                != placeholder
            ):
                continue

            box_left = shape.left
            box_top = shape.top
            box_width = shape.width
            box_height = shape.height

            with Image.open(image_file) as image:
                image_width, image_height = image.size

            scale = min(
                box_width / image_width,
                box_height / image_height
            )
            # 자리표시자 크기에 꽉 차게 맞춘다(비율 유지). 세로로 긴 범례는
            # 높이가 양식 높이와 같아진다. 위치는 자리표시자 중심을 유지한다.
            image_width = round(
                image_width * scale
            )
            image_height = round(
                image_height * scale
            )
            left = box_left + (
                box_width - image_width
            ) // 2
            top = box_top + (
                box_height - image_height
            ) // 2

            slide.shapes._spTree.remove(
                shape.element
            )

            slide.shapes.add_picture(
                image_file,
                left,
                top,
                width=image_width,
                height=image_height
            )

            break

    # =====================================
    # 결과 슬라이드 생성
    # =====================================

    def add_case_slide(
        self,
        case_name,
        rack,
        glass,
        legend_values=None
    ):

        """결과 템플릿을 복제해 Case 제목·Rack/Glass 이미지·범위 값을 채운다."""

        page_number = (
            self.base_slide_count
            + self.generated_case_count
            + 1
        )
        self.generated_case_count += 1

        slide = self.duplicate_slide(
            self.result_template_index
        )

        legend_values = legend_values or {}

        for shape in slide.shapes:

            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()

                        if config.TITLE_TAG in cell_text:
                            self.set_cell_text(
                                cell,
                                cell_text.replace(
                                    config.TITLE_TAG,
                                    case_name
                                )
                            )
                        elif row_index == 24 and column_index in (2, 26):
                            value_key = "rack" if column_index == 2 else "glass"
                            self.set_text_preserving_style(
                                cell.text_frame,
                                self.format_legend_values(
                                    legend_values.get(value_key, {})
                                )
                            )

                continue

            if not hasattr(
                shape,
                "text"
            ):
                continue

            text = (
                shape.text.strip()
            )

            # -------------------
            # Case 제목
            # -------------------

            if (
                config.TITLE_TAG
                in text
            ):

                self.set_text_preserving_style(
                    shape.text_frame,
                    text.replace(
                        config.TITLE_TAG,
                        case_name
                    )
                )

            # -------------------
            # 결과 템플릿 태그 삭제
            # -------------------

            elif (
                text
                == config.RESULT_TEMPLATE_TAG
            ):

                shape.text = ""

            elif shape.name == "Page":

                self.set_text_preserving_style(
                    shape.text_frame,
                    str(page_number)
                )

        # -----------------------
        # Rack
        # -----------------------

        if rack:

            self.replace_placeholder_image(
                slide,
                "[RACK_MODEL]",
                rack["model_file"]
            )

            self.replace_placeholder_image(
                slide,
                "[RACK_LEGEND]",
                rack["legend_file"]
            )

        # -----------------------
        # Glass
        # -----------------------

        if glass:

            self.replace_placeholder_image(
                slide,
                "[GLASS_MODEL]",
                glass["model_file"]
            )

            self.replace_placeholder_image(
                slide,
                "[GLASS_LEGEND]",
                glass["legend_file"]
            )

    @staticmethod
    def format_legend_values(values):

        """Max/Min 값을 '(-) 최소 mm ~ (+) 최대 mm' 한 줄 문자열로 만든다."""

        maximum = str(values.get("max", "")).strip()
        minimum = str(values.get("min", "")).strip()
        if maximum and minimum:
            maximum = re.sub(r"\s*mm\s*$", "", maximum, flags=re.IGNORECASE)
            minimum = re.sub(r"\s*mm\s*$", "", minimum, flags=re.IGNORECASE)
            maximum = maximum.lstrip("+- ")
            minimum = minimum.lstrip("+- ")
            return f"(-) {minimum} mm ~ (+) {maximum} mm"
        if maximum:
            return f"Max: {maximum}"
        if minimum:
            return f"Min: {minimum}"
        return ""

    # =====================================
    # 정리 슬라이드 생성
    # =====================================

    @staticmethod
    def find_summary_table_shape(slide):

        """정리 슬라이드에서 '구분/변위' 헤더와 Case 행을 가진 데이터 표를 찾는다."""

        return next(
            shape
            for shape in slide.shapes
            if getattr(shape, "has_table", False)
            and {"구분", "변위"}.issubset(
                {
                    cell.text.strip()
                    for row in shape.table.rows
                    for cell in row.cells
                    if cell.text.strip()
                }
            )
            and any(
                re.match(r"Case\s*\d+", cell.text.strip(), re.IGNORECASE)
                for row in shape.table.rows
                for cell in row.cells
            )
        )

    def add_case_summary_slide(
        self,
        case_names,
        legend_values
    ):

        """템플릿의 정리 표를 복제해 Case별 Rack/Glass 변위를 채운다."""

        if not self.summary_template_indices:
            raise ValueError("report_template.pptx에서 정리 표 양식을 찾을 수 없습니다.")

        # 마지막 정리 예시 슬라이드를 양식으로 사용한다.
        source_index = self.summary_template_indices[-1]
        template_slide = self.prs.slides[source_index]
        template_table = self.find_summary_table_shape(template_slide).table

        # 표 구조는 '헤더 1행 + Case당 Rack/Glass 2행'이므로
        # 데이터 행 수의 절반이 한 페이지에 들어가는 Case 수이다.
        capacity = (len(template_table.rows) - 1) // 2
        if capacity < 1:
            raise ValueError("정리 표 양식에 Case 데이터 행이 없습니다.")

        case_pages = [
            case_names[start:start + capacity]
            for start in range(0, len(case_names), capacity)
        ] or [[]]

        for summary_page_index, page_cases in enumerate(case_pages):
            slide = self.duplicate_slide(source_index)
            page_number = (
                self.base_slide_count
                + self.generated_case_count
                + summary_page_index
                + 1
            )
            summary_table = self.find_summary_table_shape(slide).table

            for case_offset in range(capacity):
                base_row = 1 + case_offset * 2
                if case_offset < len(page_cases):
                    case_name = page_cases[case_offset]
                    values = legend_values.get(case_name, {})
                    # 변위 칸에는 '(-) 최소 mm ~ (+) 최대 mm'를 한 칸에 합쳐 적는다.
                    rack_text = (
                        self.format_legend_values(values.get("rack", {})) or "-"
                    )
                    glass_text = (
                        self.format_legend_values(values.get("glass", {})) or "-"
                    )
                else:
                    # 양식에 남아 있는 예시 Case 값은 모두 지워 빈 행으로 둔다.
                    case_name = ""
                    rack_text = ""
                    glass_text = ""

                # 구분/Glass 두께/비고 칸은 Rack·Glass 두 행이 세로 병합되어 있어
                # 기준 셀(Rack 행)에만 쓰면 Glass 행에도 함께 표시된다.
                # Glass 두께와 비고는 사용자가 PPT에서 직접 채우도록 빈 칸으로 둔다.
                self.set_text_preserving_style(
                    summary_table.cell(base_row, 0).text_frame,
                    case_name
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row, 1).text_frame,
                    ""
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row, 2).text_frame,
                    "Rack" if case_name else ""
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row, 3).text_frame,
                    rack_text
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row + 1, 2).text_frame,
                    "Glass" if case_name else ""
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row + 1, 3).text_frame,
                    glass_text
                )
                self.set_text_preserving_style(
                    summary_table.cell(base_row, 4).text_frame,
                    ""
                )

            # 슬라이드 번호만 갱신하고 '▶ 정리' 제목 등 나머지 양식은 유지한다.
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                if shape.name == "Page" or "スライド番号" in shape.name:
                    self.set_text_preserving_style(
                        shape.text_frame,
                        str(page_number)
                    )

    # =====================================
    # 완료보고서 정보 반영
    # =====================================

    def fill_report_info(
        self,
        report_data
    ):

        """표지(제목·해석 번호 등)와 개요(안건·배경·목표·결론)를 채운다."""

        title = report_data.get(
            "title"
        )

        if report_data.get("skip_cover"):
            # 표지를 만들지 않을 때 템플릿의 안내 문구가 남지 않도록 지운다.
            for shape in self.prs.slides[0].shapes:
                if (
                    getattr(shape, "has_text_frame", False)
                    and shape.text.strip() == "보고서 제목"
                ):
                    self.set_text_preserving_style(
                        shape.text_frame,
                        ""
                    )
                    break

        elif title:
            for shape in self.prs.slides[0].shapes:
                if (
                    getattr(shape, "has_text_frame", False)
                    and shape.text.strip() == "보고서 제목"
                ):
                    self.set_text_preserving_style(
                        shape.text_frame,
                        title
                    )
                    break

        for slide_index in range(
            min(2, len(self.prs.slides))
        ):
            slide = self.prs.slides[
                slide_index
            ]

            for shape in slide.shapes:
                if not getattr(shape, "has_table", False):
                    continue

                if slide_index == 0:
                    self.fill_cover_table(
                        shape.table,
                        report_data
                    )
                elif slide_index == 1:
                    self.fill_summary_table(
                        shape.table,
                        report_data
                    )

    @staticmethod
    def set_cell_text(
        cell,
        value
    ):

        """값이 있을 때만 표 셀의 글자를 서식 유지 상태로 바꾼다."""

        if value is not None and str(value).strip():
            PPTBuilder.set_text_preserving_style(
                cell.text_frame,
                str(value)
            )

    @staticmethod
    def set_text_preserving_style(
        text_frame,
        value
    ):

        """글자만 바꾸고 템플릿의 글꼴·크기·색 등 서식은 그대로 유지한다.

        첫 번째 run의 텍스트만 교체하고 나머지 run은 제거한다.
        run이 없는 빈 단락이면 문단 끝 서식(endParaRPr)을 복사해 새 run에 적용한다.
        """

        paragraph = text_frame.paragraphs[0]

        if paragraph.runs:
            paragraph.runs[0].text = value

            for run in paragraph.runs[1:]:
                run._r.getparent().remove(run._r)
        else:
            slide_number_field = paragraph._p.find(
                qn("a:fld")
            )

            if slide_number_field is not None:
                field_text = slide_number_field.find(
                    qn("a:t")
                )

                if field_text is not None:
                    field_text.text = value
                    return

            run = paragraph.add_run()
            run.text = value

            end_properties = paragraph._p.find(
                qn("a:endParaRPr")
            )

            if end_properties is not None:
                run_properties = deepcopy(
                    end_properties
                )
                run_properties.tag = qn(
                    "a:rPr"
                )
                run._r.insert(
                    0,
                    run_properties
                )

        text_frame.word_wrap = True
        text_frame.auto_size = (
            MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        )

    def fill_cover_table(
        self,
        table,
        report_data
    ):

        """표지 표의 '해석 번호/수행 기간/담당자' 행을 찾아 값을 채운다."""

        for row in table.rows:
            label = row.cells[0].text.strip()

            if label == "해석 번호":
                self.set_cell_text(
                    row.cells[1],
                    report_data.get("analysis_number")
                )

            elif label == "수행 기간":
                self.set_cell_text(
                    row.cells[1],
                    report_data.get("start_date")
                )
                self.set_cell_text(
                    row.cells[3],
                    report_data.get("end_date")
                )

            elif label == "담당자":
                self.set_cell_text(
                    row.cells[1],
                    report_data.get("contact")
                )

    def fill_summary_table(
        self,
        table,
        report_data
    ):

        """개요 표에서 '안건명/배경/목표/최종 결론' 라벨 아래 행에 내용을 채운다."""

        field_names = {
            "안건명": "agenda",
            "배경": "background",
            "목표": "objective",
            "최종 결론": "conclusion"
        }

        for row_index, row in enumerate(table.rows):
            label = (
                row.cells[0].text
                .replace("|", "")
                .strip()
            )
            field_name = field_names.get(
                label
            )

            if (
                field_name
                and row_index + 1 < len(table.rows)
            ):
                self.set_cell_text(
                    table.rows[row_index + 1].cells[0],
                    report_data.get(field_name)
                )

    # =====================================
    # 슬라이드 삭제
    # =====================================

    def remove_slide(
        self,
        slide_index
    ):

        """슬라이드 목록과 관계(rel)에서 해당 슬라이드를 완전히 제거한다."""

        slide_id = (
            self.prs.slides._sldIdLst[
                slide_index
            ]
        )

        self.prs.part.drop_rel(
            slide_id.rId
        )

        del self.prs.slides._sldIdLst[
            slide_index
        ]

    # =====================================
    # 저장
    # =====================================

    def save(self, output_path=None):

        """편집용 템플릿 슬라이드를 제거한 뒤 최종 PPT를 저장한다.

        output_path가 없으면 config의 기본 출력 경로에 저장한다.
        """

        template_indices = {
            self.result_template_index,
            *self.summary_template_indices
        }
        # 뒤에서부터 지워야 앞 슬라이드 번호가 밀리지 않는다.
        for slide_index in sorted(template_indices, reverse=True):
            self.remove_slide(slide_index)

        self.prs.save(
            str(
                output_path or config.OUTPUT_PPT
            )
        )