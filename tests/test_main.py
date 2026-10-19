import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np
from PIL import Image
from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from src.detectors.legend_detector import LegendDetector
from src.detectors.model_detector import ModelDetector
from src.gui import format_report_text, parse_report_text
from main import (
    build_cases,
    build_cases_from_assignments,
    build_legend_values,
    generate_presentation
)
from src.workflow import parse_image_metadata, rename_assigned_images
from src.processors.report_reader import read_report_data
from src.ppt.ppt_builder import PPTBuilder


class BuildCasesTest(unittest.TestCase):
    def test_build_cases_accepts_case_variations(self):
        files = [
            "C:/tmp/CaseA_Rack.png",
            "C:/tmp/CaseA_Glass.png",
            "C:/tmp/case_b_rack.jpg",
            "C:/tmp/case_b_glass.jpeg",
            "C:/tmp/other.png",
        ]

        cases = build_cases(files)

        self.assertEqual(set(cases.keys()), {"CaseA", "case_b"})
        self.assertIn("rack", cases["CaseA"])
        self.assertIn("glass", cases["CaseA"])
        self.assertIn("rack", cases["case_b"])
        self.assertIn("glass", cases["case_b"])

    def test_build_cases_from_manual_assignments(self):
        cases = build_cases_from_assignments({
            "Case 2": {"Glass": "image-a.png"},
            "Case 1": {"rack": "image-b.png", "unknown": "image-c.png"},
            "   ": {"glass": "image-d.png"}
        })

        self.assertEqual(
            cases,
            {
                "Case 2": {"glass": "image-a.png"},
                "Case 1": {"rack": "image-b.png"}
            }
        )

    def test_build_legend_values_groups_by_case_and_part(self):
        self.assertEqual(
            build_legend_values(
                {
                    "glass.png": ("Case A", "glass"),
                    "rack.png": ("Case A", "rack")
                },
                {
                    "glass.png": {"max": "0.34", "min": "-3.74"},
                    "rack.png": {"max": "0.25", "min": "-3.75"}
                }
            ),
            {
                "Case A": {
                    "glass": {"max": "0.34", "min": "-3.74"},
                    "rack": {"max": "0.25", "min": "-3.75"}
                }
            }
        )

    def test_parse_image_metadata_reads_generated_case_and_extrema(self):
        metadata = parse_image_metadata(
            "C:/results/Case 3_Glass_Max_0.33986_Min_-3.7426.png"
        )

        self.assertEqual(metadata, {
            "case": "Case 3",
            "part": "glass",
            "max": "0.33986",
            "min": "-3.7426"
        })

    def test_rename_assigned_images_preserves_collisions_and_metadata(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "original.png"
            expected_name = "Case 2_Glass_Max_0.34_Min_-3.74.png"
            existing_target = Path(temp_dir) / expected_name
            source.write_bytes(b"source image")
            existing_target.write_bytes(b"existing file")

            renamed, errors = rename_assigned_images(
                {"Case 2": {"glass": str(source)}},
                {"Case 2": {"glass": {"max": "0.34", "min": "-3.74"}}}
            )

            renamed_path = Path(renamed[str(source)])
            self.assertEqual(renamed_path.name, "Case 2_Glass_Max_0.34_Min_-3.74_2.png")
            self.assertFalse(source.exists())
            self.assertEqual(existing_target.read_bytes(), b"existing file")
            self.assertEqual(renamed_path.read_bytes(), b"source image")
            self.assertEqual(
                parse_image_metadata(renamed_path),
                {"case": "Case 2", "part": "glass", "max": "0.34", "min": "-3.74"}
            )
            self.assertEqual(errors, [])

    def test_report_text_editor_round_trips_all_report_fields(self):
        report = {
            "title": "Glass 처짐",
            "analysis_number": "CAE-001",
            "contact": "홍길동",
            "start_date": "2026-10-01",
            "end_date": "2026-10-19",
            "agenda": "Half Glass 평가",
            "background": "하중 조건 확인\n기존 설계 비교",
            "objective": "최대 처짐 검토",
            "conclusion": "Case 3 적용 검토"
        }

        parsed = parse_report_text(format_report_text(report))

        for field, value in report.items():
            self.assertEqual(parsed[field], value)


class PresentationWorkflowTest(unittest.TestCase):
    def test_generation_processes_selected_report_and_case_images(self):
        workbook = Workbook()
        workbook.active.title = "해석 완료 보고서"
        workbook.create_sheet("해석 결과 보고서")

        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "report.xlsx"
            output_path = Path(temp_dir) / "result.pptx"
            workbook.save(report_path)
            logs = []

            with patch("src.workflow.config.OUTPUT_PPT", output_path), \
                    patch("src.workflow.ImageProcessor") as image_processor, \
                    patch("src.workflow.PPTBuilder") as ppt_builder:
                image_processor.return_value.process.return_value = {
                    "model_file": "model.png",
                    "legend_file": "legend.png"
                }

                result_path, case_count = generate_presentation(
                    str(report_path),
                    [
                        "C:/results/image_01.png",
                        "C:/results/plot-final.png",
                        "C:/results/view_left.png"
                    ],
                    progress_callback=logs.append,
                    case_order=["CaseB", "CaseA"],
                    report_overrides={
                        "title": "GUI에서 수정한 제목",
                        "analysis_number": "GUI-001",
                        "contact": "작성자",
                        "start_date": "2026-08-01",
                        "end_date": "2026-08-09",
                        "agenda": "안건 입력",
                        "background": "배경 입력",
                        "objective": "목표 입력",
                        "conclusion": "최종 결론 입력"
                    },
                    case_assignments={
                        "CaseA": {
                            "rack": "C:/results/image_01.png",
                            "glass": "C:/results/plot-final.png"
                        },
                        "CaseB": {
                            "rack": "C:/results/view_left.png"
                        }
                    },
                    legend_values={
                        "CaseA": {
                            "glass": {"max": "0.34", "min": "-3.74"}
                        }
                    }
                )

        self.assertEqual(result_path, output_path)
        self.assertEqual(case_count, 2)
        self.assertEqual(image_processor.return_value.process.call_count, 3)
        self.assertEqual(ppt_builder.return_value.add_case_slide.call_count, 2)
        self.assertEqual(
            [call.args[0] for call in ppt_builder.return_value.add_case_slide.call_args_list],
            ["CaseB", "CaseA"]
        )
        passed_report_data = ppt_builder.return_value.fill_report_info.call_args.args[0]
        self.assertEqual(passed_report_data["title"], "GUI에서 수정한 제목")
        self.assertEqual(passed_report_data["analysis_number"], "GUI-001")
        self.assertEqual(passed_report_data["contact"], "작성자")
        self.assertEqual(passed_report_data["start_date"], "2026-08-01")
        self.assertEqual(passed_report_data["end_date"], "2026-08-09")
        self.assertEqual(passed_report_data["agenda"], "안건 입력")
        self.assertEqual(passed_report_data["background"], "배경 입력")
        self.assertEqual(passed_report_data["objective"], "목표 입력")
        self.assertEqual(passed_report_data["conclusion"], "최종 결론 입력")
        self.assertEqual(
            ppt_builder.return_value.add_case_slide.call_args_list[1]
            .kwargs["legend_values"],
            {"glass": {"max": "0.34", "min": "-3.74"}}
        )
        self.assertEqual(
            ppt_builder.return_value.add_case_slide.call_args_list[0]
            .kwargs["legend_values"],
            {}
        )
        ppt_builder.return_value.save.assert_called_once_with(None)
        self.assertIn("PowerPoint 파일을 저장하는 중입니다.", logs)

    def test_generation_can_skip_report_and_cover(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "result.pptx"

            with patch("src.workflow.config.OUTPUT_PPT", output_path), \
                    patch("src.workflow.read_report_data") as read_report, \
                    patch("src.workflow.ImageProcessor") as image_processor, \
                    patch("src.workflow.PPTBuilder") as ppt_builder:
                image_processor.return_value.process.return_value = {
                    "model_file": "model.png",
                    "legend_file": "legend.png"
                }

                result_path, case_count = generate_presentation(
                    "",
                    ["C:/results/plot_01.png"],
                    case_assignments={
                        "Manual Case": {
                            "glass": "C:/results/plot_01.png"
                        }
                    }
                )

        self.assertEqual(result_path, output_path)
        self.assertEqual(case_count, 1)
        read_report.assert_not_called()
        self.assertTrue(
            ppt_builder.return_value.fill_report_info
            .call_args.args[0]["skip_cover"]
        )

    def test_generation_uses_explicit_output_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            default_path = Path(temp_dir) / "default.pptx"
            custom_path = Path(temp_dir) / "report_2026.pptx"

            with patch("src.workflow.config.OUTPUT_PPT", default_path), \
                    patch("src.workflow.ImageProcessor") as image_processor, \
                    patch("src.workflow.PPTBuilder") as ppt_builder:
                image_processor.return_value.process.return_value = {
                    "model_file": "model.png",
                    "legend_file": "legend.png"
                }

                result_path, case_count = generate_presentation(
                    "",
                    ["C:/results/plot_01.png"],
                    output_path=custom_path,
                    case_assignments={
                        "Manual Case": {
                            "glass": "C:/results/plot_01.png"
                        }
                    }
                )

        # 지정한 저장 경로가 저장 호출과 반환값에 그대로 쓰여야 한다.
        self.assertEqual(result_path, custom_path)
        self.assertEqual(case_count, 1)
        ppt_builder.return_value.save.assert_called_once_with(custom_path)


class ReportImportTest(unittest.TestCase):
    def test_generated_tables_inherit_template_cell_font_style(self):
        builder = PPTBuilder()
        style_slide = builder.prs.slides[
            builder.summary_template_indices[-1]
        ]
        style_table = next(
            shape.table
            for shape in style_slide.shapes
            if getattr(shape, "has_table", False)
            and any(
                cell.text.strip() == "구분"
                for row in shape.table.rows
                for cell in row.cells
            )
        )
        header_source = next(
            cell
            for row in style_table.rows
            for cell in row.cells
            if cell.text.strip() == "구분"
        )
        header_source.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

        builder.add_case_summary_slide(["Case 1"], {})

        # 마지막 슬라이드가 방금 생성된 정리 슬라이드이며 템플릿 글꼴을 그대로 물려받는다.
        generated_slide = builder.prs.slides[len(builder.prs.slides) - 1]
        generated_table = PPTBuilder.find_summary_table_shape(
            generated_slide
        ).table
        self.assertEqual(
            generated_table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size.pt,
            16
        )

    def test_saved_presentation_ends_with_case_deflection_summary(self):
        builder = PPTBuilder()
        legend_values = {
            "Case 1": {
                "rack": {"max": "0.0176", "min": "-3.29"},
                "glass": {"max": "0.0113", "min": "-2.1"}
            }
        }
        builder.add_case_slide(
            "Case 1",
            None,
            None,
            legend_values=legend_values["Case 1"]
        )
        builder.add_case_summary_slide(
            ["Case 1"],
            legend_values
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "summary.pptx"
            with patch("src.ppt.ppt_builder.config.OUTPUT_PPT", output_path):
                builder.save()

            presentation = Presentation(output_path)
            last_slide = presentation.slides[len(presentation.slides) - 1]
            summary_table = PPTBuilder.find_summary_table_shape(
                last_slide
            ).table
            page_text = next(
                shape.text.strip()
                for shape in last_slide.shapes
                if getattr(shape, "has_text_frame", False)
                and (
                    shape.name == "Page"
                    or "スライド番号" in shape.name
                )
            )

        self.assertEqual(summary_table.cell(0, 0).text, "구분")
        self.assertEqual(summary_table.cell(0, 1).text, "Glass 두께")
        self.assertEqual(summary_table.cell(0, 2).text, "변위")
        self.assertEqual(summary_table.cell(0, 4).text, "비고")
        self.assertEqual(summary_table.cell(1, 0).text, "Case 1")
        self.assertEqual(summary_table.cell(1, 1).text, "")
        self.assertEqual(summary_table.cell(1, 2).text, "Rack")
        self.assertEqual(
            summary_table.cell(1, 3).text,
            "(-) 3.29 mm ~ (+) 0.0176 mm"
        )
        self.assertEqual(summary_table.cell(2, 2).text, "Glass")
        self.assertEqual(
            summary_table.cell(2, 3).text,
            "(-) 2.1 mm ~ (+) 0.0113 mm"
        )
        self.assertEqual(summary_table.cell(1, 4).text, "")
        # 사용하지 않은 양식 행의 예시 값은 비워져 있어야 한다.
        self.assertEqual(summary_table.cell(3, 0).text, "")
        self.assertEqual(page_text, str(builder.base_slide_count + 2))

    def test_skip_cover_clears_default_title_placeholder(self):
        builder = PPTBuilder()
        builder.fill_report_info({"skip_cover": True})

        title_text = next(
            shape.text
            for shape in builder.prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
        )

        self.assertEqual(title_text, "")

    def test_template_uses_report_style_headings_and_keeps_placeholders(self):
        builder = PPTBuilder()
        expected_headings = {
            1: "해석 개요 | Summary of Analysis",
            2: "해석 수행 내용 및 성과 | 해석 모델 및 조건 (1/1)",
            3: "해석 수행 내용 및 성과 | 해석 결과"
        }

        for slide_index, expected_heading in expected_headings.items():
            slide = builder.prs.slides[slide_index]
            heading = next(
                shape
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
                and expected_heading in shape.text
            )
            run = heading.text_frame.paragraphs[0].runs[0]

            self.assertEqual(run.font.size.pt, 20)
            self.assertTrue(run.font.bold)

        result_slide = builder.prs.slides[
            builder.result_template_index
        ]
        table_cells = [
            cell
            for shape in result_slide.shapes
            if getattr(shape, "has_table", False)
            for row in shape.table.rows
            for cell in row.cells
        ]
        placeholders = {
            cell.text.strip()
            for cell in table_cells
        }
        placeholders.update(
            shape.text.strip()
            for shape in result_slide.shapes
            if getattr(shape, "has_text_frame", False)
        )

        self.assertTrue(
            any("[CASE_TITLE]" in cell.text for cell in table_cells)
        )
        self.assertIn("[RACK_MODEL]", placeholders)
        self.assertIn("[RACK_LEGEND]", placeholders)
        self.assertIn("[GLASS_MODEL]", placeholders)
        self.assertIn("[GLASS_LEGEND]", placeholders)

    def test_reads_report_fields_from_xlsx(self):
        workbook = Workbook()
        report_sheet = workbook.active
        report_sheet.title = "해석 완료 보고서"
        workbook.create_sheet("해석 결과 보고서")

        report_sheet["B5"] = "해석완료보고서"
        report_sheet["F2"] = "CAE2500-123"
        report_sheet["T2"] = "홍길동"
        report_sheet["AD2"] = date(2026, 4, 30)
        report_sheet["J23"] = "안건 내용"
        report_sheet["J30"] = "배경 내용"
        report_sheet["J38"] = "목표 내용"
        report_sheet["J46"] = "결론 내용"

        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = os.path.join(temp_dir, "report.xlsx")
            workbook.save(file_path)

            report_data = read_report_data(file_path)

        self.assertEqual(report_data["title"], "해석완료보고서")
        self.assertEqual(report_data["analysis_number"], "CAE2500-123")
        self.assertEqual(report_data["contact"], "홍길동")
        # 수행 종료일은 '해석 완료 보고서' 시트 AD2에서 읽는다.
        self.assertEqual(report_data["end_date"], "2026-04-30")
        self.assertNotIn("start_date", report_data)
        self.assertEqual(report_data["agenda"], "안건 내용")
        self.assertEqual(report_data["background"], "배경 내용")
        self.assertEqual(report_data["objective"], "목표 내용")
        self.assertEqual(report_data["conclusion"], "결론 내용")

    def test_fills_cover_and_summary_slides(self):
        builder = PPTBuilder()
        title_shape = next(
            shape
            for shape in builder.prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
            and shape.text.strip() == "보고서 제목"
        )
        original_title_size = (
            title_shape.text_frame.paragraphs[0]
            .runs[0].font.size.pt
        )
        cover_table = next(
            shape.table
            for shape in builder.prs.slides[0].shapes
            if getattr(shape, "has_table", False)
        )
        original_cell_margin = (
            cover_table.rows[1].cells[1].margin_left
        )
        conclusion = "최종 결론 내용입니다. " * 35

        builder.fill_report_info({
            "title": "완료보고서",
            "analysis_number": "CAE2500-123",
            "contact": "홍길동",
            "start_date": "2026-04-01",
            "end_date": "2026-04-30",
            "agenda": "안건 내용",
            "background": "배경 내용",
            "objective": "목표 내용",
            "conclusion": conclusion
        })

        title = next(
            shape.text
            for shape in builder.prs.slides[0].shapes
            if getattr(shape, "has_text_frame", False)
            and shape.text.strip() == "완료보고서"
        )
        summary_table = next(
            shape.table
            for shape in builder.prs.slides[1].shapes
            if getattr(shape, "has_table", False)
        )

        self.assertEqual(title, "완료보고서")
        self.assertEqual(
            title_shape.text_frame.paragraphs[0].runs[0].font.size.pt,
            original_title_size
        )
        self.assertEqual(
            title_shape.text_frame.auto_size,
            MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        )
        self.assertEqual(cover_table.rows[1].cells[1].text, "CAE2500-123")
        self.assertEqual(
            cover_table.rows[1].cells[1]
            .text_frame.paragraphs[0].runs[0].font.size.pt,
            11
        )
        self.assertEqual(
            cover_table.rows[1].cells[1].margin_left,
            original_cell_margin
        )
        self.assertEqual(cover_table.rows[2].cells[1].text, "2026-04-01")
        self.assertEqual(cover_table.rows[2].cells[3].text, "2026-04-30")
        self.assertEqual(cover_table.rows[3].cells[1].text, "홍길동")
        self.assertEqual(summary_table.rows[1].cells[0].text, "안건 내용")
        self.assertEqual(summary_table.rows[3].cells[0].text, "배경 내용")
        self.assertEqual(summary_table.rows[5].cells[0].text, "목표 내용")
        self.assertEqual(
            summary_table.rows[5].cells[0]
            .text_frame.paragraphs[0].runs[0].font.size.pt,
            12
        )
        self.assertEqual(
            summary_table.rows[5].cells[0]
            .text_frame.paragraphs[0].runs[0].font.name,
            "맑은 고딕"
        )
        self.assertEqual(summary_table.rows[7].cells[0].text, conclusion)
        self.assertEqual(
            summary_table.rows[7].cells[0]
            .text_frame.paragraphs[0].runs[0].font.size.pt,
            12
        )
        self.assertTrue(summary_table.rows[7].cells[0].text_frame.word_wrap)
        self.assertEqual(
            summary_table.rows[7].cells[0].text_frame.auto_size,
            MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        )

    def test_result_slides_keep_case_titles_and_sequential_pages(self):
        builder = PPTBuilder()
        first_result_index = len(builder.prs.slides)
        builder.add_case_slide("Case 1", None, None)
        builder.add_case_slide("Case 2", None, None)

        for slide_index, expected_title, expected_page in [
            (first_result_index, "Case 1", str(builder.base_slide_count + 1)),
            (first_result_index + 1, "Case 2", str(builder.base_slide_count + 2))
        ]:
            slide = builder.prs.slides[slide_index]
            title_cell = next(
                cell
                for shape in slide.shapes
                if getattr(shape, "has_table", False)
                for row in shape.table.rows
                for cell in row.cells
                if expected_title in cell.text
            )
            page = next(
                shape
                for shape in slide.shapes
                if shape.name == "Page"
            )

            self.assertTrue(
                title_cell.text.strip().endswith(expected_title)
            )
            self.assertEqual(page.text.strip(), expected_page)

    def test_case_slide_writes_rack_and_glass_legend_extrema(self):
        builder = PPTBuilder()
        result_slide_index = len(builder.prs.slides)
        builder.add_case_slide(
            "Case 1",
            None,
            None,
            legend_values={
                "rack": {"max": "0.25", "min": "-3.75"},
                "glass": {"max": "0.34", "min": "-3.74"}
            }
        )
        slide = builder.prs.slides[result_slide_index]
        table = next(
            shape.table
            for shape in slide.shapes
            if getattr(shape, "has_table", False)
        )

        self.assertEqual(
            table.rows[24].cells[2].text,
            "(-) 3.75 mm ~ (+) 0.25 mm"
        )
        self.assertEqual(
            table.rows[24].cells[26].text,
            "(-) 3.74 mm ~ (+) 0.34 mm"
        )


class ImagePlacementTest(unittest.TestCase):
    def test_images_are_contained_and_centered_in_placeholders(self):
        for image_size in [
            (1600, 900),
            (800, 1200),
            (600, 1200),
            (1000, 1000)
        ]:
            with self.subTest(image_size=image_size):
                builder = PPTBuilder()
                slide = builder.prs.slides[
                    builder.result_template_index
                ]
                placeholder = next(
                    shape
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False)
                    and shape.text.strip() == "[RACK_MODEL]"
                )
                box = (
                    placeholder.left,
                    placeholder.top,
                    placeholder.width,
                    placeholder.height
                )

                with tempfile.TemporaryDirectory() as temp_dir:
                    image_path = os.path.join(
                        temp_dir,
                        "model.png"
                    )
                    Image.new(
                        "RGB",
                        image_size,
                        color="white"
                    ).save(image_path)

                    builder.replace_placeholder_image(
                        slide,
                        "[RACK_MODEL]",
                        image_path
                    )

                picture = next(
                    shape
                    for shape in slide.shapes
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                )

                self.assertLessEqual(picture.width, box[2])
                self.assertLessEqual(picture.height, box[3])
                self.assertAlmostEqual(
                    picture.left + picture.width / 2,
                    box[0] + box[2] / 2,
                    delta=1
                )
                self.assertAlmostEqual(
                    picture.top + picture.height / 2,
                    box[1] + box[3] / 2,
                    delta=1
                )
                self.assertAlmostEqual(
                    picture.width / picture.height,
                    image_size[0] / image_size[1],
                    delta=0.002
                )

    def test_legend_image_fills_placeholder_with_contain_fit(self):
        builder = PPTBuilder()
        slide = builder.prs.slides[builder.result_template_index]
        placeholder = next(
            shape
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            and shape.text.strip() == "[RACK_LEGEND]"
        )
        box = (
            placeholder.left,
            placeholder.top,
            placeholder.width,
            placeholder.height
        )
        source_size = (600, 900)

        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = os.path.join(temp_dir, "legend.png")
            Image.new("RGB", source_size, color="white").save(image_path)
            builder.replace_placeholder_image(
                slide,
                "[RACK_LEGEND]",
                image_path
            )

        picture = next(
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        # 범례도 축소 없이 자리표시자에 꽉 차게(contain) 들어가야 한다.
        contain_scale = min(
            box[2] / source_size[0],
            box[3] / source_size[1]
        )
        expected_width = round(source_size[0] * contain_scale)
        expected_height = round(source_size[1] * contain_scale)
        self.assertAlmostEqual(picture.width, expected_width, delta=1)
        self.assertAlmostEqual(picture.height, expected_height, delta=1)
        self.assertAlmostEqual(
            picture.left + picture.width / 2,
            box[0] + box[2] / 2,
            delta=1
        )


class LegendDetectorTest(unittest.TestCase):
    def test_detects_legend_in_broader_search_region(self):
        detector = LegendDetector()
        img = np.full((800, 1200, 3), 255, dtype=np.uint8)
        cv2.rectangle(img, (640, 180), (820, 620), (0, 0, 0), -1)

        with tempfile.TemporaryDirectory() as temp_dir:
            result = detector.detect(img, "legend_case", temp_dir)
            self.assertTrue(result)
            self.assertTrue(os.path.exists(result))


class ModelDetectorTest(unittest.TestCase):
    def test_prefers_tall_centered_model_over_legend_block(self):
        detector = ModelDetector()
        img = np.full((700, 1000, 3), 255, dtype=np.uint8)

        cv2.rectangle(img, (0, 0), (260, 260), (0, 0, 0), -1)
        cv2.rectangle(img, (300, 100), (520, 620), (120, 120, 120), -1)
        cv2.rectangle(img, (300, 120), (520, 600), (255, 0, 0), -1)
        cv2.rectangle(img, (70, 520), (180, 650), (0, 0, 0), -1)

        with tempfile.TemporaryDirectory() as temp_dir:
            result = detector.detect(img, "model_case", temp_dir)
            self.assertTrue(result)
            self.assertTrue(os.path.exists(result))


if __name__ == "__main__":
    unittest.main()
